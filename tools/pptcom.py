"""pptcom: a python-pptx stand-in for the PowerPoint COM object model, so the engine (deckkit, the family pattern
libraries, states) runs where PowerPoint does not: Linux, the claude.ai sandbox, CI.

Only the COM subset the engine uses is implemented. Coordinates are points, colours are COM RGB integers
(R + 256 G + 65536 B), indexes are 1-based, exactly as in COM. Auto-size text boxes are measured with the
design-system fonts' real advance widths (fontmetrics.json, calibrated against PowerPoint: widths within 0.3 %,
line pitch 1.2117 em x line spacing), so layouts that stack on measured heights match the PowerPoint build.

PDF export (SaveAs ..., 32) uses LibreOffice (soffice) when present, else PowerPoint (Windows), else fails loudly.
"""
import copy, json, os, re, shutil, subprocess, tempfile
from lxml import etree
from pptx import Presentation as _Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

E = 12700                       # EMU per point
TRUE, FALSE = -1, 0
HERE = os.path.dirname(os.path.abspath(__file__))
_METRICS = json.load(open(os.path.join(HERE, 'fontmetrics.json'), encoding='utf-8'))
PITCH = 1.2117                  # PowerPoint single line pitch, in em (all design fonts)
_FIRST = [(0.0, 0.0), (1.0, 1.2117), (1.05, 1.2117), (1.2, 1.3193), (1.5, 1.5919), (2.0, 2.0463), (3.0, 2.95)]
_FILLS = ('noFill', 'solidFill', 'gradFill', 'blipFill', 'pattFill', 'grpFill')
_ORDER = {
    'rPr': ['ln', *_FILLS, 'effectLst', 'effectDag', 'highlight', 'uLnTx', 'uLn', 'uFillTx', 'uFill', 'latin', 'ea',
            'cs', 'sym', 'hlinkClick', 'hlinkMouseOver', 'rtl', 'extLst'],
    'pPr': ['lnSpc', 'spcBef', 'spcAft', 'buClrTx', 'buClr', 'buSzTx', 'buSzPct', 'buSzPts', 'buFontTx', 'buFont',
            'buNone', 'buAutoNum', 'buChar', 'buBlip', 'tabLst', 'defRPr', 'extLst'],
    'spPr': ['xfrm', 'custGeom', 'prstGeom', *_FILLS, 'ln', 'effectLst', 'effectDag', 'scene3d', 'sp3d', 'extLst'],
    'ln': [*_FILLS, 'prstDash', 'custDash', 'round', 'bevel', 'miter', 'headEnd', 'tailEnd', 'extLst'],
}
_BU = ('buNone', 'buAutoNum', 'buChar', 'buBlip')


def _local(el):
    return etree.QName(el).localname


def _child(parent, tag, create=True, exclusive=()):
    """Child a:<tag> of parent, created at its schema position; `exclusive` siblings are removed first."""
    for t in exclusive:
        if t != tag:
            for c in parent.findall(qn('a:' + t)):
                parent.remove(c)
    c = parent.find(qn('a:' + tag))
    if c is not None or not create:
        return c
    c = etree.SubElement(parent, qn('a:' + tag))
    order = _ORDER.get(_local(parent))
    if order and tag in order:
        parent.remove(c)
        rank = order.index(tag)
        idx = len(parent)
        for i, sib in enumerate(parent):
            n = _local(sib)
            if n in order and order.index(n) > rank:
                idx = i
                break
        parent.insert(idx, c)
    return c


def _hex(v):
    v = int(v)
    return '%02X%02X%02X' % (v & 255, (v >> 8) & 255, (v >> 16) & 255)


def _int(h):
    return int(h[0:2], 16) + int(h[2:4], 16) * 256 + int(h[4:6], 16) * 65536


def _solid(parent, v, exclusive=_FILLS):
    f = _child(parent, 'solidFill', exclusive=exclusive)
    for c in list(f):
        f.remove(c)
    etree.SubElement(f, qn('a:srgbClr'), val=_hex(v))
    return f


# ================================================================ text measurement
def _font_key(name):
    if name in _METRICS:
        return name
    base = re.sub(r' (Semibold|SemiBold|Light|Bold)$', '', name or '')
    return base if base in _METRICS else 'Segoe UI'


def _font_m(name, bold):
    return _METRICS[_font_key(name)]['bold' if bold else 'regular']


def _char_w(ch, name, size, bold):
    """Advance width in points. A glyph the font lacks takes the width PowerPoint's font fallback gives it
    (measured, 'fallback' table), else Cambria Math's, else Segoe UI Symbol's, else the font's average."""
    cp = str(ord(ch))
    m = _font_m(name, bold)
    w = m['w'].get(cp)
    if w is None:
        w = _METRICS.get('fallback', {}).get(_font_key(name), {}).get(cp)
    if w is None:
        for fb in ('Cambria Math', 'Segoe UI Symbol'):
            w = _METRICS.get(fb, {}).get('regular', {}).get('w', {}).get(cp)
            if w is not None:
                break
    return (w if w is not None else m['avg']) * size


def _first_line(s):
    for (a, ha), (b, hb) in zip(_FIRST, _FIRST[1:]):
        if s <= b:
            return ha + (hb - ha) * (s - a) / (b - a)
    return _FIRST[-1][1]


def _run_style(r, default):
    rpr = r.find(qn('a:rPr'))
    d = dict(default)
    if rpr is not None:
        if rpr.get('sz'):
            d['size'] = int(rpr.get('sz')) / 100
        if rpr.get('b') is not None:
            d['bold'] = rpr.get('b') in ('1', 'true')
        lat = rpr.find(qn('a:latin'))
        if lat is not None and not lat.get('typeface', '+').startswith('+'):
            d['font'] = lat.get('typeface')
    return d


def _para_lines(p, width, default):
    """Greedy PowerPoint-like wrap of one paragraph: -> [(line_width, max_size)] (breaks after spaces/hyphens)."""
    chars = []
    for el in p:
        n = _local(el)
        if n in ('r', 'fld'):
            st = _run_style(el, default)
            for ch in el.findtext(qn('a:t')) or '':
                chars.append((ch, _char_w(ch, st['font'], st['size'], st['bold']), st['size']))
        elif n == 'br':
            chars.append(('\n', 0, _run_style(el, default)['size']))
    epr = p.find(qn('a:endParaRPr'))
    end_size = int(epr.get('sz')) / 100 if epr is not None and epr.get('sz') else default['size']
    if not chars:
        return [(0, end_size)]
    lines, cur, cur_w, brk = [], [], 0.0, -1
    for ch, w, sz in chars:
        if ch == '\n':
            lines.append(cur); cur, cur_w, brk = [], 0.0, -1
            continue
        if width and cur and cur_w + w > width + 0.01 and not ch.isspace():
            if brk >= 0:
                lines.append(cur[:brk + 1]); cur = cur[brk + 1:]
            else:
                lines.append(cur); cur = []
            cur_w = sum(c[1] for c in cur); brk = -1
        cur.append((ch, w, sz)); cur_w += w
        if ch.isspace() or ch in '-‐–—':
            brk = len(cur) - 1
    lines.append(cur)
    out = []
    for ln in lines:
        vis = ln[:]
        while vis and vis[-1][0].isspace():
            vis.pop()
        out.append((sum(c[1] for c in vis), max([c[2] for c in ln] or [end_size])))
    return out


def measure(txBody, width=None):
    """(height, widest line) in points of a text body laid out at `width` (None = no wrap)."""
    bp = txBody.find(qn('a:bodyPr'))
    ins = lambda k, d: int(bp.get(k)) / E if bp is not None and bp.get(k) is not None else d
    l, r, t, b = ins('lIns', 7.2), ins('rIns', 7.2), ins('tIns', 3.6), ins('bIns', 3.6)
    default = dict(font='Segoe UI', size=18.0, bold=False)
    h, widest, first = 0.0, 0.0, True
    for p in txBody.findall(qn('a:p')):
        ppr = p.find(qn('a:pPr'))
        s, before, after, marl, ind = 1.0, 0.0, 0.0, 0.0, 0.0
        if ppr is not None:
            ls = ppr.find(qn('a:lnSpc'))
            if ls is not None and ls.find(qn('a:spcPct')) is not None:
                s = int(ls.find(qn('a:spcPct')).get('val')) / 100000
            for tag, key in (('spcBef', 'before'), ('spcAft', 'after')):
                e = ppr.find(qn('a:' + tag))
                if e is not None and e.find(qn('a:spcPts')) is not None:
                    v = int(e.find(qn('a:spcPts')).get('val')) / 100
                    before, after = (v, after) if key == 'before' else (before, v)
            marl = int(ppr.get('marL', 0)) / E
            ind = int(ppr.get('indent', 0)) / E
        avail = None if width is None else max(1.0, width - l - r - marl)
        lines = _para_lines(p, avail, default)
        if not first:
            h += before + prev_after
        for k, (lw, sz) in enumerate(lines):
            h += (_first_line(s) if (first and k == 0) else PITCH * s) * sz
            widest = max(widest, lw + marl + (ind if k == 0 else 0))
        first, prev_after = False, after
    return h + t + b, widest + l + r


# ================================================================ text ranges
class _Color:
    def __init__(self, get, set_):
        self._get, self._set = get, set_

    @property
    def RGB(self):
        return self._get()

    @RGB.setter
    def RGB(self, v):
        self._set(v)


class _Font:
    def __init__(self, rng):
        self._r = rng

    def _rprs(self):
        return self._r._rprs()

    def _first(self):
        rs = self._rprs()
        return rs[0] if rs else None

    @property
    def Name(self):
        f = self._first()
        lat = f.find(qn('a:latin')) if f is not None else None
        return lat.get('typeface') if lat is not None else ''

    @Name.setter
    def Name(self, v):
        for rpr in self._rprs():
            _child(rpr, 'latin').set('typeface', v)

    @property
    def Size(self):
        f = self._first()
        return int(f.get('sz')) / 100 if f is not None and f.get('sz') else 18

    @Size.setter
    def Size(self, v):
        for rpr in self._rprs():
            rpr.set('sz', str(int(round(v * 100))))

    def _flag(self, attr, v):
        for rpr in self._rprs():
            rpr.set(attr, '1' if v and v != FALSE else '0')

    @property
    def Bold(self):
        f = self._first()
        return TRUE if f is not None and f.get('b') == '1' else FALSE

    @Bold.setter
    def Bold(self, v):
        self._flag('b', v)

    @property
    def Italic(self):
        f = self._first()
        return TRUE if f is not None and f.get('i') == '1' else FALSE

    @Italic.setter
    def Italic(self, v):
        self._flag('i', v)

    @property
    def Subscript(self):
        f = self._first()
        return TRUE if f is not None and int(f.get('baseline', 0)) < 0 else FALSE

    @Subscript.setter
    def Subscript(self, v):
        for rpr in self._rprs():
            if v and v != FALSE:
                rpr.set('baseline', '-25000')
            elif 'baseline' in rpr.attrib:
                del rpr.attrib['baseline']

    @property
    def Color(self):
        def get():
            f = self._first()
            c = f.find(qn('a:solidFill') + '/' + qn('a:srgbClr')) if f is not None else None
            return _int(c.get('val')) if c is not None else 0

        def set_(v):
            for rpr in self._rprs():
                _solid(rpr, v)
        return _Color(get, set_)


class _Bullet:
    def __init__(self, pprs):
        self._pprs = pprs

    @property
    def Visible(self):
        p = self._pprs()
        return TRUE if p and p[0].find(qn('a:buChar')) is not None else FALSE

    @Visible.setter
    def Visible(self, v):
        for ppr in self._pprs():
            if v and v != FALSE:
                if ppr.find(qn('a:buChar')) is None:
                    _child(ppr, 'buChar', exclusive=_BU).set('char', '•')
            else:
                _child(ppr, 'buNone', exclusive=_BU)

    @property
    def Character(self):
        p = self._pprs()
        c = p[0].find(qn('a:buChar')) if p else None
        return ord(c.get('char')) if c is not None else 8226

    @Character.setter
    def Character(self, v):
        for ppr in self._pprs():
            _child(ppr, 'buChar', exclusive=_BU).set('char', chr(v))

    @property
    def Font(self):
        pprs = self._pprs

        class F:
            @property
            def Color(self_):
                def set_(v):
                    for ppr in pprs():
                        c = _child(ppr, 'buClr')
                        for x in list(c):
                            c.remove(x)
                        etree.SubElement(c, qn('a:srgbClr'), val=_hex(v))
                return _Color(lambda: 0, set_)
        return F()


class _ParagraphFormat:
    AL = {1: 'l', 2: 'ctr', 3: 'r', 4: 'just'}

    def __init__(self, rng):
        self._r = rng

    def _pprs(self):
        out = []
        for p in self._r._paras():
            ppr = p.find(qn('a:pPr'))
            if ppr is None:
                ppr = etree.Element(qn('a:pPr'))
                p.insert(0, ppr)
            out.append(ppr)
        return out

    @property
    def Alignment(self):
        p = self._pprs()
        inv = {v: k for k, v in self.AL.items()}
        return inv.get(p[0].get('algn', 'l'), 1) if p else 1

    @Alignment.setter
    def Alignment(self, v):
        for ppr in self._pprs():
            ppr.set('algn', self.AL[v])

    def _spacing(self, tag, v, pct):
        for ppr in self._pprs():
            e = _child(ppr, tag)
            for c in list(e):
                e.remove(c)
            if pct:
                etree.SubElement(e, qn('a:spcPct'), val=str(int(round(v * 100000))))
            else:
                etree.SubElement(e, qn('a:spcPts'), val=str(int(round(v * 100))))

    SpaceWithin = property(lambda s: 1.0, lambda s, v: s._spacing('lnSpc', v, True))
    SpaceAfter = property(lambda s: 0.0, lambda s, v: s._spacing('spcAft', v, False))
    SpaceBefore = property(lambda s: 0.0, lambda s, v: s._spacing('spcBef', v, False))

    @property
    def Bullet(self):
        return _Bullet(self._pprs)


class TextRange:
    """A character range of a text body. Paragraph breaks count as one character ('\\r'), as in COM."""

    def __init__(self, frame, start=0, length=None):
        self._f, self._s, self._n = frame, start, length

    # -- flattening
    def _body(self):
        return self._f._body(create=True)

    def _layout(self):
        """[(p, [(run_el, start, end)], p_start, p_end)] over the whole body."""
        out, pos = [], 0
        for p in self._body().findall(qn('a:p')):
            runs, p0 = [], pos
            for el in p:
                n = _local(el)
                if n in ('r', 'fld'):
                    t = el.findtext(qn('a:t')) or ''
                    runs.append((el, pos, pos + len(t))); pos += len(t)
                elif n == 'br':
                    runs.append((el, pos, pos + 1)); pos += 1
            out.append((p, runs, p0, pos))
            pos += 1                                   # the paragraph break
        return out

    def _total(self):
        lay = self._layout()
        return lay[-1][3] if lay else 0

    def _span(self):
        end = self._total() if self._n is None else min(self._s + self._n, self._total())
        return self._s, end

    def _split_at(self, pos):
        for p, runs, p0, p1 in self._layout():
            for el, a, b in runs:
                if a < pos < b and _local(el) == 'r':
                    t = el.find(qn('a:t'))
                    left, right = t.text[:pos - a], t.text[pos - a:]
                    new = copy.deepcopy(el)
                    t.text = left
                    new.find(qn('a:t')).text = right
                    el.addnext(new)
                    return

    def _runs(self):
        a, b = self._span()
        self._split_at(a); self._split_at(b)
        out = []
        for p, runs, p0, p1 in self._layout():
            for el, s, e in runs:
                if _local(el) == 'r' and s >= a and e <= b and e > s:
                    out.append(el)
        return out

    def _rprs(self):
        out = []
        for el in self._runs():
            rpr = el.find(qn('a:rPr'))
            if rpr is None:
                rpr = etree.Element(qn('a:rPr'), lang='en-US')
                el.insert(0, rpr)
            out.append(rpr)
        a, b = self._span()
        for p, runs, p0, p1 in self._layout():         # paragraph end marks inside the range follow the range
            if a <= p1 <= b and (p0 < b or not runs):
                epr = p.find(qn('a:endParaRPr'))
                if epr is None:
                    epr = etree.SubElement(p, qn('a:endParaRPr'), lang='en-US')
                out.append(epr)
        return out

    def _paras(self):
        a, b = self._span()
        return [p for p, runs, p0, p1 in self._layout() if p0 <= b and p1 >= a]

    # -- COM surface
    @property
    def Text(self):
        body = self._body()
        paras = []
        for p in body.findall(qn('a:p')):
            s = ''
            for el in p:
                n = _local(el)
                if n in ('r', 'fld'):
                    s += el.findtext(qn('a:t')) or ''
                elif n == 'br':
                    s += '\v'
            paras.append(s)
        full = '\r'.join(paras)
        a, b = self._span()
        return full[a:b]

    @Text.setter
    def Text(self, s):
        body = self._body()
        paras = body.findall(qn('a:p'))
        whole = self._s == 0 and (self._n is None or self._n >= self._total())
        if not whole:
            runs = self._runs()
            if runs:
                runs[0].find(qn('a:t')).text = s
                for el in runs[1:]:
                    el.getparent().remove(el)
            return
        p0 = paras[0] if paras else None
        ppr = copy.deepcopy(p0.find(qn('a:pPr'))) if p0 is not None and p0.find(qn('a:pPr')) is not None else None
        r0 = p0.find(qn('a:r')) if p0 is not None else None
        rpr = copy.deepcopy(r0.find(qn('a:rPr'))) if r0 is not None and r0.find(qn('a:rPr')) is not None else None
        if rpr is None and p0 is not None and p0.find(qn('a:endParaRPr')) is not None:
            rpr = copy.deepcopy(p0.find(qn('a:endParaRPr')))
            rpr.tag = qn('a:rPr')
        epr = copy.deepcopy(p0.find(qn('a:endParaRPr'))) if p0 is not None and p0.find(qn('a:endParaRPr')) is not None else None
        for p in paras:
            body.remove(p)
        for line in s.replace('\n', '\r').split('\r'):
            p = etree.SubElement(body, qn('a:p'))
            if ppr is not None:
                p.append(copy.deepcopy(ppr))
            for k, seg in enumerate(line.split('\v')):
                if k:
                    br = etree.SubElement(p, qn('a:br'))
                    if rpr is not None:
                        br.append(copy.deepcopy(rpr))
                if seg:
                    r = etree.SubElement(p, qn('a:r'))
                    r.append(copy.deepcopy(rpr) if rpr is not None else etree.Element(qn('a:rPr'), lang='en-US'))
                    etree.SubElement(r, qn('a:t')).text = seg
            if epr is not None:
                p.append(copy.deepcopy(epr))
            elif rpr is not None:
                e = copy.deepcopy(rpr); e.tag = qn('a:endParaRPr'); p.append(e)

    @property
    def Font(self):
        return _Font(self)

    @property
    def ParagraphFormat(self):
        return _ParagraphFormat(self)

    def Characters(self, start, length):
        return TextRange(self._f, self._s + start - 1, length)

    def Runs(self, i=None):
        a, b = self._span()
        spans = [(s, e) for p, runs, p0, p1 in self._layout() for el, s, e in runs
                 if _local(el) == 'r' and s >= a and e <= b]
        if i is None:
            class Rs:
                Count = len(spans)
            return Rs()
        s, e = spans[i - 1]
        return TextRange(self._f, s, e - s)

    def Replace(self, find, repl):
        hit = None
        for p, runs, p0, p1 in self._layout():
            rs = [(el, s, e) for el, s, e in runs if _local(el) == 'r']
            text = ''.join(el.findtext(qn('a:t')) or '' for el, s, e in rs)
            k = text.find(find)
            if k < 0:
                continue
            a, b = p0 + k, p0 + k + len(find)
            inside = [x for x in rs if x[1] < b and x[2] > a]
            first = inside[0]
            t0 = first[0].find(qn('a:t'))
            last = inside[-1]
            tail = (last[0].findtext(qn('a:t')) or '')[b - last[1]:]
            t0.text = (t0.text or '')[:a - first[1]] + repl + tail
            for el, s, e in inside[1:]:
                el.getparent().remove(el)
            hit = self
            break
        return hit

    @property
    def BoundHeight(self):
        return measure(self._body(), self._f._wrap_width())[0]

    @property
    def BoundWidth(self):
        return measure(self._body(), self._f._wrap_width())[1]


class _Level:
    def __init__(self, frame):
        self._f = frame

    def _pprs(self):
        return _ParagraphFormat(self._f.TextRange)._pprs()

    @property
    def FirstMargin(self):
        p = self._pprs()[0]
        return (int(p.get('marL', 0)) + int(p.get('indent', 0))) / E

    @FirstMargin.setter
    def FirstMargin(self, v):
        for p in self._pprs():
            p.set('indent', str(int(round(v * E)) - int(p.get('marL', 0))))

    @property
    def LeftMargin(self):
        return int(self._pprs()[0].get('marL', 0)) / E

    @LeftMargin.setter
    def LeftMargin(self, v):
        for p in self._pprs():
            first = int(p.get('marL', 0)) + int(p.get('indent', 0))
            p.set('marL', str(int(round(v * E))))
            p.set('indent', str(first - int(round(v * E))))


class TextFrame:
    AN = {1: 't', 3: 'ctr', 4: 'b'}

    def __init__(self, shape):
        self._sh = shape

    def _body(self, create=False):
        el = self._sh._el
        body = el.find(qn('p:txBody'))
        if body is None and create:
            body = etree.SubElement(el, qn('p:txBody'))
            etree.SubElement(body, qn('a:bodyPr'))
            etree.SubElement(body, qn('a:lstStyle'))
            etree.SubElement(body, qn('a:p'))
        return body

    def _bp(self):
        return self._body(create=True).find(qn('a:bodyPr'))

    def _ins(self, k, v=None):
        if v is None:
            x = self._bp().get(k)
            return int(x) / E if x is not None else (7.2 if k in ('lIns', 'rIns') else 3.6)
        self._bp().set(k, str(int(round(v * E))))

    MarginLeft = property(lambda s: s._ins('lIns'), lambda s, v: s._ins('lIns', v))
    MarginRight = property(lambda s: s._ins('rIns'), lambda s, v: s._ins('rIns', v))
    MarginTop = property(lambda s: s._ins('tIns'), lambda s, v: s._ins('tIns', v))
    MarginBottom = property(lambda s: s._ins('bIns', ), lambda s, v: s._ins('bIns', v))

    @property
    def AutoSize(self):
        return 1 if self._bp().find(qn('a:spAutoFit')) is not None else 0

    @AutoSize.setter
    def AutoSize(self, v):
        bp = self._bp()
        for t in ('noAutofit', 'normAutofit', 'spAutoFit'):
            for c in bp.findall(qn('a:' + t)):
                bp.remove(c)
        c = etree.Element(qn('a:spAutoFit' if v == 1 else 'a:noAutofit'))
        warp = bp.find(qn('a:prstTxWarp'))
        bp.insert(1 if warp is not None else 0, c)
        if v == 1:
            self._sh._fit()

    @property
    def WordWrap(self):
        return FALSE if self._bp().get('wrap') == 'none' else TRUE

    @WordWrap.setter
    def WordWrap(self, v):
        self._bp().set('wrap', 'square' if v and v != FALSE else 'none')

    @property
    def VerticalAnchor(self):
        return {v: k for k, v in self.AN.items()}.get(self._bp().get('anchor', 't'), 1)

    @VerticalAnchor.setter
    def VerticalAnchor(self, v):
        self._bp().set('anchor', self.AN[v])

    @property
    def TextRange(self):
        return TextRange(self)

    @property
    def HasText(self):
        b = self._body()
        return TRUE if b is not None and ''.join(t.text or '' for t in b.iter(qn('a:t'))) else FALSE

    @property
    def Ruler(self):
        frame = self

        class Ruler:
            def Levels(self_, i):
                return _Level(frame)
        return Ruler()

    def _wrap_width(self):
        return None if self._bp().get('wrap') == 'none' else self._sh.Width


# ================================================================ shapes
class _Fill:
    def __init__(self, sh):
        self._sh = sh

    def _sppr(self):
        return self._sh._sppr()

    @property
    def Visible(self):
        return FALSE if self._sppr().find(qn('a:noFill')) is not None else TRUE

    @Visible.setter
    def Visible(self, v):
        if not v or v == FALSE:
            _child(self._sppr(), 'noFill', exclusive=_FILLS)
        elif self._sppr().find(qn('a:noFill')) is not None:
            self._sppr().remove(self._sppr().find(qn('a:noFill')))

    def Solid(self):
        sp = self._sppr()
        if sp.find(qn('a:solidFill')) is None:
            _solid(sp, 0xFFFFFF)

    def _clr(self):
        f = self._sppr().find(qn('a:solidFill'))
        return f.find(qn('a:srgbClr')) if f is not None else None

    @property
    def ForeColor(self):
        def get():
            c = self._clr()
            return _int(c.get('val')) if c is not None else None

        def set_(v):
            old = self.Transparency if self._clr() is not None else 0
            _solid(self._sppr(), v)
            if old:
                self.Transparency = old
        return _Color(get, set_)

    @property
    def Transparency(self):
        c = self._clr()
        a = c.find(qn('a:alpha')) if c is not None else None
        return 1 - int(a.get('val')) / 100000 if a is not None else 0.0

    @Transparency.setter
    def Transparency(self, v):
        c = self._clr()
        if c is None:
            self.Solid(); c = self._clr()
        for a in c.findall(qn('a:alpha')):
            c.remove(a)
        if v:
            etree.SubElement(c, qn('a:alpha'), val=str(int(round((1 - v) * 100000))))


class _Line:
    STYLE = {1: 'none', 2: 'triangle', 3: 'arrow', 4: 'stealth', 5: 'diamond', 6: 'oval'}
    SIZE = {1: 'sm', 2: 'med', 3: 'lg'}
    DASH = {1: 'solid', 2: 'sysDot', 3: 'sysDot', 4: 'dash', 5: 'dashDot', 6: 'lgDashDotDot', 7: 'lgDash',
            8: 'lgDashDot'}

    def __init__(self, sh):
        self._sh = sh

    def _ln(self):
        return _child(self._sh._sppr(), 'ln')

    @property
    def Visible(self):
        ln = self._sh._sppr().find(qn('a:ln'))
        return FALSE if ln is not None and ln.find(qn('a:noFill')) is not None else TRUE

    @Visible.setter
    def Visible(self, v):
        ln = self._ln()
        if not v or v == FALSE:
            _child(ln, 'noFill', exclusive=_FILLS)
        elif ln.find(qn('a:noFill')) is not None:
            ln.remove(ln.find(qn('a:noFill')))

    @property
    def ForeColor(self):
        def get():
            c = self._ln().find(qn('a:solidFill') + '/' + qn('a:srgbClr'))
            return _int(c.get('val')) if c is not None else None
        return _Color(get, lambda v: _solid(self._ln(), v))

    @property
    def Weight(self):
        return int(self._ln().get('w', 12700)) / E

    @Weight.setter
    def Weight(self, v):
        self._ln().set('w', str(int(round(v * E))))

    def _end(self, which, attr, v):
        e = _child(self._ln(), which)
        e.set(attr, v)
        if attr != 'type' and e.get('type') is None:
            e.set('type', 'none')

    EndArrowheadStyle = property(lambda s: 1, lambda s, v: s._end('tailEnd', 'type', s.STYLE[v]))
    EndArrowheadLength = property(lambda s: 2, lambda s, v: s._end('tailEnd', 'len', s.SIZE[v]))
    EndArrowheadWidth = property(lambda s: 2, lambda s, v: s._end('tailEnd', 'w', s.SIZE[v]))
    BeginArrowheadStyle = property(lambda s: 1, lambda s, v: s._end('headEnd', 'type', s.STYLE[v]))
    BeginArrowheadLength = property(lambda s: 2, lambda s, v: s._end('headEnd', 'len', s.SIZE[v]))
    BeginArrowheadWidth = property(lambda s: 2, lambda s, v: s._end('headEnd', 'w', s.SIZE[v]))

    @property
    def DashStyle(self):
        return 1

    @DashStyle.setter
    def DashStyle(self, v):
        _child(self._ln(), 'prstDash').set('val', self.DASH.get(v, 'solid'))
        if v == 3:
            self._ln().set('cap', 'rnd')


class _Shadow:
    def __init__(self, sh):
        self._sh = sh

    @property
    def Visible(self):
        return TRUE

    @Visible.setter
    def Visible(self, v):
        if not v or v == FALSE:
            sp = self._sh._sppr()
            if sp.find(qn('a:effectLst')) is None:
                _child(sp, 'effectLst')


class _Adjustments:
    def __init__(self, sh):
        self._sh = sh

    def SetItem(self, i, v):
        self._sh._pp.adjustments[i - 1] = v

    def Item(self, i):
        return self._sh._pp.adjustments[i - 1]


class _PlaceholderFormat:
    def __init__(self, el):
        self._el = el

    @property
    def Type(self):
        ph = self._el.find('.//' + qn('p:ph'))
        t = ph.get('type', 'obj') if ph is not None else 'obj'
        return {'title': 1, 'body': 2, 'ctrTitle': 3, 'subTitle': 4, 'dt': 16, 'sldNum': 13, 'ftr': 15,
                'obj': 7, 'pic': 18, 'tbl': 12, 'chart': 8}.get(t, 7)


class Shape:
    """COM-like view of one slide shape element (sp, cxnSp, pic, grpSp, graphicFrame)."""

    def __init__(self, slide, el, pp_shape=None):
        self._slide, self._el, self._ppx = slide, el, pp_shape

    @property
    def _pp(self):
        if self._ppx is None:
            from pptx.shapes.shapetree import SlideShapeFactory
            self._ppx = SlideShapeFactory(self._el, self._slide._pp.shapes)
        return self._ppx

    def _nv(self):
        return self._el.find('.//' + qn('p:cNvPr'))

    def _sppr(self):
        sp = self._el.find(qn('p:spPr'))
        if sp is None:
            sp = self._el.find(qn('p:grpSpPr'))
        return sp

    def _xfrm(self, create=False):
        sp = self._sppr()
        x = sp.find(qn('a:xfrm')) if sp is not None else None
        if x is None and create:
            x = _child(sp, 'xfrm')
            etree.SubElement(x, qn('a:off'), x='0', y='0')
            etree.SubElement(x, qn('a:ext'), cx='0', cy='0')
        return x

    def _geo(self, i):
        x = self._xfrm()
        if x is None:                                  # placeholder: inherit from the layout
            try:
                return [self._pp.left, self._pp.top, self._pp.width, self._pp.height][i] / E
            except Exception:
                return 0.0
        off, ext = x.find(qn('a:off')), x.find(qn('a:ext'))
        return int([off.get('x'), off.get('y'), ext.get('cx'), ext.get('cy')][i]) / E

    def _set_geo(self, i, v):
        if self._xfrm() is None:
            cur = [self._geo(k) for k in range(4)]
            x = self._xfrm(create=True)
            x.find(qn('a:off')).set('x', str(int(cur[0] * E))); x.find(qn('a:off')).set('y', str(int(cur[1] * E)))
            x.find(qn('a:ext')).set('cx', str(int(cur[2] * E))); x.find(qn('a:ext')).set('cy', str(int(cur[3] * E)))
        x = self._xfrm()
        el, attr = [(x.find(qn('a:off')), 'x'), (x.find(qn('a:off')), 'y'), (x.find(qn('a:ext')), 'cx'),
                    (x.find(qn('a:ext')), 'cy')][i]
        el.set(attr, str(int(round(v * E))))

    Left = property(lambda s: s._geo(0), lambda s, v: s._set_geo(0, v))
    Top = property(lambda s: s._geo(1), lambda s, v: s._set_geo(1, v))
    Width = property(lambda s: s._geo(2), lambda s, v: (s._set_geo(2, v), s._fit()))

    @property
    def Height(self):
        self._fit()
        return self._geo(3)

    @Height.setter
    def Height(self, v):
        self._set_geo(3, v)

    def _fit(self):
        """Auto-size (shape to fit text): recompute the height (and the width when not wrapping)."""
        body = self._el.find(qn('p:txBody'))
        if body is None or body.find(qn('a:bodyPr') + '/' + qn('a:spAutoFit')) is None or self._xfrm() is None:
            return
        bp = body.find(qn('a:bodyPr'))
        nowrap = bp.get('wrap') == 'none'
        h, w = measure(body, None if nowrap else self._geo(2))
        if nowrap:
            left, width = self._geo(0), self._geo(2)
            algn = (body.find(qn('a:p') + '/' + qn('a:pPr')).get('algn', 'l')
                    if body.find(qn('a:p') + '/' + qn('a:pPr')) is not None else 'l')
            if algn == 'ctr':
                self._set_geo(0, left + (width - w) / 2)
            elif algn == 'r':
                self._set_geo(0, left + width - w)
            self._set_geo(2, w)
        anchor = bp.get('anchor', 't')
        top, old = self._geo(1), self._geo(3)
        if anchor == 'b':
            self._set_geo(1, top + old - h)
        elif anchor == 'ctr':
            self._set_geo(1, top + (old - h) / 2)
        self._set_geo(3, h)

    @property
    def Name(self):
        return self._nv().get('name', '')

    @Name.setter
    def Name(self, v):
        self._nv().set('name', v)

    @property
    def Id(self):
        return int(self._nv().get('id'))

    @property
    def Type(self):
        n = _local(self._el)
        if n == 'grpSp':
            return 6
        if n == 'cxnSp':
            return 9
        if n == 'pic':
            return 13
        if n == 'graphicFrame':
            return 7
        if self._el.find('.//' + qn('p:ph')) is not None:
            return 14
        sp = self._sppr()
        if sp is not None and sp.find(qn('a:custGeom')) is not None:
            return 5
        if self._el.find(qn('p:nvSpPr') + '/' + qn('p:cNvSpPr')) is not None and \
                self._el.find(qn('p:nvSpPr') + '/' + qn('p:cNvSpPr')).get('txBox') == '1':
            return 17
        return 1

    @property
    def HasTextFrame(self):
        return TRUE if _local(self._el) == 'sp' else FALSE

    @property
    def TextFrame(self):
        return TextFrame(self)

    TextFrame2 = TextFrame

    @property
    def Fill(self):
        return _Fill(self)

    @property
    def Line(self):
        return _Line(self)

    @property
    def Shadow(self):
        return _Shadow(self)

    @property
    def Adjustments(self):
        return _Adjustments(self)

    @property
    def PlaceholderFormat(self):
        return _PlaceholderFormat(self._el)

    @property
    def GroupItems(self):
        kids = [c for c in self._el if _local(c) in ('sp', 'cxnSp', 'pic', 'grpSp', 'graphicFrame')]
        slide = self._slide

        class G:
            Count = len(kids)

            def __call__(self_, i):
                return Shape(slide, kids[i - 1])

            def __iter__(self_):
                return (Shape(slide, k) for k in kids)
        return G()

    def ZOrder(self, cmd):
        tree = self._el.getparent()
        tree.remove(self._el)
        if cmd in (0, 2):                              # msoBringToFront / msoBringForward
            tree.append(self._el)
        else:                                          # msoSendToBack / msoSendBackward
            tree.insert(2, self._el)

    def Delete(self):
        self._el.getparent().remove(self._el)

    def __eq__(self, other):
        return isinstance(other, Shape) and other._el is self._el

    def __hash__(self):
        return id(self._el)


_SHAPE_TAGS = ('sp', 'cxnSp', 'pic', 'grpSp', 'graphicFrame', 'contentPart')


class _FreeformBuilder:
    """COM BuildFreeform: AddNodes(segment 0=line / 1=curve, editing, x1, y1[, x2, y2, x3, y3]) -> ConvertToShape."""

    def __init__(self, shapes, x, y):
        self._shapes, self._cmds = shapes, [('M', [(x, y)])]

    def AddNodes(self, seg, edit, x1, y1, x2=None, y2=None, x3=None, y3=None):
        if seg == 1 and x3 is not None:
            self._cmds.append(('C', [(x1, y1), (x2, y2), (x3, y3)]))
        else:
            self._cmds.append(('L', [(x1, y1)]))

    def ConvertToShape(self):
        pts = [p for _, ps in self._cmds for p in ps]
        x0, y0 = min(p[0] for p in pts), min(p[1] for p in pts)
        x1, y1 = max(p[0] for p in pts), max(p[1] for p in pts)
        w, h = max(x1 - x0, 0.01), max(y1 - y0, 0.01)
        shp = self._shapes.AddShape(1, x0, y0, w, h)
        sp = shp._sppr()
        prst = sp.find(qn('a:prstGeom'))
        cust = etree.Element(qn('a:custGeom'))
        for t in ('avLst', 'gdLst', 'ahLst', 'cxnLst'):
            etree.SubElement(cust, qn('a:' + t))
        etree.SubElement(cust, qn('a:rect'), l='l', t='t', r='r', b='b')
        pl = etree.SubElement(etree.SubElement(cust, qn('a:pathLst')), qn('a:path'),
                              w=str(int(round(w * E))), h=str(int(round(h * E))))
        rel = lambda p: (str(int(round((p[0] - x0) * E))), str(int(round((p[1] - y0) * E))))
        tags = {'M': 'moveTo', 'L': 'lnTo', 'C': 'cubicBezTo'}
        for c, ps in self._cmds:
            e = etree.SubElement(pl, qn('a:' + tags[c]))
            for p in ps:
                x, y = rel(p)
                etree.SubElement(e, qn('a:pt'), x=x, y=y)
        first, last = self._cmds[0][1][0], self._cmds[-1][1][-1]
        if abs(first[0] - last[0]) < 0.01 and abs(first[1] - last[1]) < 0.01:
            etree.SubElement(pl, qn('a:close'))
        sp.replace(prst, cust)
        shp.Name = f'Freeform {shp.Id - 1}'
        return shp


class Shapes:
    def __init__(self, slide):
        self._slide = slide

    def _tree(self):
        return self._slide._pp.shapes._spTree

    def _els(self):
        return [c for c in self._tree() if _local(c) in _SHAPE_TAGS]

    @property
    def Count(self):
        return len(self._els())

    def __call__(self, i):
        if isinstance(i, str):
            return next(Shape(self._slide, e) for e in self._els() if Shape(self._slide, e).Name == i)
        return Shape(self._slide, self._els()[i - 1])

    Item = __call__

    def __iter__(self):
        return iter([Shape(self._slide, e) for e in self._els()])

    def __len__(self):
        return self.Count

    def _wrap(self, pp_shape):
        return Shape(self._slide, pp_shape._element, pp_shape)

    def AddTextbox(self, orientation, x, y, w, h):
        s = self._wrap(self._slide._pp.shapes.add_textbox(int(x * E), int(y * E), int(w * E), int(h * E)))
        bp = s.TextFrame._bp()
        for c in list(bp):
            bp.remove(c)
        bp.set('wrap', 'square'); bp.set('rtlCol', '0')
        return s

    def AddShape(self, kind, x, y, w, h):
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
        s = self._wrap(self._slide._pp.shapes.add_shape(MSO_AUTO_SHAPE_TYPE(kind), int(x * E), int(y * E),
                                                        int(w * E), int(h * E)))
        return s

    def AddLine(self, x1, y1, x2, y2):
        c = self._slide._pp.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(round(x1 * E)), int(round(y1 * E)),
                                                 int(round(x2 * E)), int(round(y2 * E)))
        return self._wrap(c)

    def AddPicture(self, path, link, save, x, y, w=-1, h=-1):
        kw = {}
        if w and w > 0: kw['width'] = int(w * E)
        if h and h > 0: kw['height'] = int(h * E)
        return self._wrap(self._slide._pp.shapes.add_picture(path, int(x * E), int(y * E), **kw))

    def BuildFreeform(self, edit, x, y):
        return _FreeformBuilder(self, x, y)

    @property
    def Title(self):
        for e in self._els():
            ph = e.find('.//' + qn('p:ph'))
            if ph is not None and ph.get('type') in ('title', 'ctrTitle'):
                return Shape(self._slide, e)
        return None

    @property
    def Placeholders(self):
        slide = self._slide
        els = [e for e in self._els() if e.find('.//' + qn('p:ph')) is not None]

        class P:
            Count = len(els)

            def __call__(self_, i):
                return Shape(slide, els[i - 1])

            def __iter__(self_):
                return iter([Shape(slide, e) for e in els])
        return P()


# ================================================================ slides, layouts, presentation
class Layout:
    def __init__(self, pres, pp_layout):
        self._pres, self._pp = pres, pp_layout

    @property
    def Name(self):
        return self._pp.name

    @Name.setter
    def Name(self, v):
        self._pp._element.cSld.set('name', v)

    @property
    def Shapes(self):
        return Shapes(self)


class Slide:
    def __init__(self, pres, pp_slide):
        self._pres, self._pp = pres, pp_slide

    @property
    def Shapes(self):
        return Shapes(self)

    @property
    def Name(self):
        return self._pp._element.cSld.get('name', '')

    @Name.setter
    def Name(self, v):
        self._pp._element.cSld.set('name', v)

    @property
    def SlideIndex(self):
        return self._pres._index(self._pp) + 1

    @property
    def CustomLayout(self):
        return Layout(self._pres, self._pp.slide_layout)

    def Delete(self):
        self._pres._delete(self._pp)

    def MoveTo(self, i):
        self._pres._move(self._pp, i - 1)

    def Duplicate(self):
        new = self._pres._duplicate(self._pp)

        class R:
            Count = 1

            def Item(self_, i):
                return new
        return R()

    def __eq__(self, other):
        return isinstance(other, Slide) and other._pp._element is self._pp._element

    def __hash__(self):
        return id(self._pp._element)


class _Slides:
    def __init__(self, pres):
        self._pres = pres

    @property
    def Count(self):
        return len(self._pres._pp.slides)

    def __call__(self, i):
        return Slide(self._pres, self._pres._pp.slides[i - 1])

    Item = __call__

    def __iter__(self):
        return iter([Slide(self._pres, s) for s in self._pres._pp.slides])

    def AddSlide(self, index, layout):
        s = self._pres._pp.slides.add_slide(layout._pp)
        self._pres._move(s, index - 1)
        return Slide(self._pres, s)


class _CustomProps:
    NS = 'http://schemas.openxmlformats.org/officeDocument/2006/custom-properties'
    VT = 'http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'

    def __init__(self, pres):
        self._pres = pres

    def _part(self):
        for p in self._pres._pp.part.package.iter_parts():
            if str(p.partname) == '/docProps/custom.xml':
                return p
        return None

    def _props(self):
        p = self._part()
        return (p, etree.fromstring(p.blob)) if p is not None else (None, None)

    @property
    def Count(self):
        p, x = self._props()
        return len(x) if x is not None else 0

    def __call__(self, i):
        p, x = self._props()
        el = x[i - 1]
        outer = self

        class Prop:
            Name = el.get('name')

            @property
            def Value(self_):
                return el[0].text if len(el) else None

            @Value.setter
            def Value(self_, v):
                el[0].text = str(v)
                p._blob = etree.tostring(x, xml_declaration=True, encoding='UTF-8', standalone=True)
        return Prop()

    def Add(self, name, link, typ, value):
        p, x = self._props()
        if p is None:
            return None                                  # templates built by the engine are already stamped
        pid = max([int(e.get('pid')) for e in x] + [1]) + 1
        e = etree.SubElement(x, '{%s}property' % self.NS, fmtid='{D5CDD505-2E9C-101B-9397-08002B2CF9AE}',
                             pid=str(pid), name=name)
        etree.SubElement(e, '{%s}lpwstr' % self.VT).text = str(value)
        p._blob = etree.tostring(x, xml_declaration=True, encoding='UTF-8', standalone=True)


class Presentation:
    def __init__(self, path):
        self.FullName = os.path.abspath(path)
        self._pp = _Presentation(path)

    # -- internals
    def _ids(self):
        return self._pp.slides._sldIdLst

    def _index(self, pp_slide):
        for k, s in enumerate(self._pp.slides):
            if s._element is pp_slide._element:
                return k
        raise ValueError('slide not in presentation')

    def _move(self, pp_slide, pos):
        lst = self._ids()
        el = list(lst)[self._index(pp_slide)]
        lst.remove(el)
        lst.insert(pos, el)

    def _delete(self, pp_slide):
        lst = self._ids()
        el = list(lst)[self._index(pp_slide)]
        self._pp.part.drop_rel(el.rId)
        lst.remove(el)

    def _duplicate(self, pp_slide):
        new = self._pp.slides.add_slide(pp_slide.slide_layout)
        src_tree, dst_tree = pp_slide.shapes._spTree, new.shapes._spTree
        for c in list(dst_tree):
            if _local(c) in _SHAPE_TAGS:
                dst_tree.remove(c)
        rmap = {}
        for rid, rel in pp_slide.part.rels.items():
            if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
                continue
            rmap[rid] = (new.part.relate_to(rel.target_ref, rel.reltype, is_external=True) if rel.is_external
                         else new.part.relate_to(rel.target_part, rel.reltype))
        for c in src_tree:
            if _local(c) in _SHAPE_TAGS:
                cc = copy.deepcopy(c)
                for el in cc.iter():
                    for k, v in list(el.attrib.items()):
                        if k.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}') \
                                and v in rmap:
                            el.set(k, rmap[v])
                dst_tree.append(cc)
        bg = pp_slide._element.cSld.find(qn('p:bg'))
        if bg is not None:
            new._element.cSld.insert(0, copy.deepcopy(bg))
        new._element.cSld.set('name', pp_slide._element.cSld.get('name', ''))
        self._move(new, self._index(pp_slide) + 1)
        return Slide(self, new)

    def _fit_all(self):
        for s in self._pp.slides:
            for c in s.shapes._spTree.iter(qn('p:sp')):
                Shape(Slide(self, s), c)._fit()

    # -- COM surface
    @property
    def Slides(self):
        return _Slides(self)

    @property
    def CustomDocumentProperties(self):
        return _CustomProps(self)

    def Designs(self, i):
        pres = self
        master = self._pp.slide_masters[i - 1]

        class D:
            class SlideMaster:
                pass
        sm = D.SlideMaster()

        class CL:
            @property
            def Count(self_):
                return len(master.slide_layouts)

            def __call__(self_, k):
                return Layout(pres, master.slide_layouts[k - 1])
        sm.CustomLayouts = CL()
        d = D(); d.SlideMaster = sm
        return d

    @property
    def PageSetup(self):
        pp = self._pp

        class PS:
            SlideWidth = pp.slide_width / E
            SlideHeight = pp.slide_height / E
        return PS()

    def Save(self):
        self._fit_all()
        self._pp.save(self.FullName)

    def SaveAs(self, path, fmt=None):
        path = os.path.abspath(path)
        if fmt == 32 or path.lower().endswith('.pdf'):
            self._fit_all()
            tmp = tempfile.mkdtemp()
            src = os.path.join(tmp, os.path.splitext(os.path.basename(path))[0] + '.pptx')
            self._pp.save(src)
            to_pdf(src, path)
            shutil.rmtree(tmp, ignore_errors=True)
            return
        self.FullName = path
        self.Save()

    def Close(self):
        pass


def soffice():
    for c in ('soffice', 'libreoffice', r'C:\Program Files\LibreOffice\program\soffice.exe'):
        p = shutil.which(c) or (c if os.path.exists(c) else None)
        if p:
            return p
    return None


def to_pdf(pptx, pdf):
    """PPTX -> PDF: LibreOffice if available, else PowerPoint (Windows)."""
    exe = soffice()
    if exe:
        out = tempfile.mkdtemp()
        env = dict(os.environ)
        fc = os.path.join(HERE, 'fonts.conf')
        if os.path.exists(fc) and 'FONTCONFIG_FILE' not in env:
            env['FONTCONFIG_FILE'] = fc
        subprocess.run([exe, '--headless', '--convert-to', 'pdf', '--outdir', out, pptx], check=True,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, env=env, timeout=600)
        produced = os.path.join(out, os.path.splitext(os.path.basename(pptx))[0] + '.pdf')
        shutil.move(produced, pdf)
        shutil.rmtree(out, ignore_errors=True)
        return pdf
    try:
        import win32com.client
    except ImportError:
        raise SystemExit('PDF export needs LibreOffice (soffice) or PowerPoint; neither was found.')
    app = win32com.client.Dispatch('PowerPoint.Application')
    pres = app.Presentations.Open(os.path.abspath(pptx), -1, 0, 0)
    try:
        pres.SaveAs(os.path.abspath(pdf), 32)
    finally:
        pres.Close()
    return pdf


class Application:
    """COM 'PowerPoint.Application' look-alike."""

    class _Presentations:
        def Open(self, path, read_only=0, untitled=0, window=0):
            return Presentation(path)

    Presentations = _Presentations()

    def Quit(self):
        pass
