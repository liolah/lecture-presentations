"""List the shapes on slides of a deck (name, type, box in points, text snippet, picture pixel size).

Usage:  python tools/shapes.py DECK.pptx [--slides 4,12-16]
Used when surveying originals, e.g. to find the figure boxes before cropping them with figcrop.py.
"""
import argparse, io, sys
from pptx import Presentation
from pptx.util import Emu

E = 12700


def sel(s, n):
    if not s:
        return range(1, n + 1)
    out = set()
    for p in s.split(','):
        a, _, b = p.partition('-'); out.update(range(int(a), int(b or a) + 1))
    return sorted(out)


def main():
    ap = argparse.ArgumentParser(); ap.add_argument('deck'); ap.add_argument('--slides')
    a = ap.parse_args()
    prs = Presentation(a.deck)
    print(f'canvas {prs.slide_width / E:.0f} x {prs.slide_height / E:.0f} pt')
    for i in sel(a.slides, len(prs.slides)):
        sl = prs.slides[i - 1]
        print(f'--- slide {i}')
        for sh in sl.shapes:
            if sh.width is None:
                continue
            box = f'({sh.left / E:.0f},{sh.top / E:.0f}) {sh.width / E:.0f}x{sh.height / E:.0f}'
            extra = ''
            if sh.shape_type == 13:  # picture
                try:
                    from PIL import Image
                    im = Image.open(io.BytesIO(sh.image.blob))
                    extra = f' img {im.size[0]}x{im.size[1]} {sh.image.ext}'
                except Exception as e:
                    extra = f' img ? ({sh.image.ext if hasattr(sh, "image") else e})'
            descr = sh._element.xpath('./*[1]/p:cNvPr/@descr')
            if descr:
                extra += f' alt={descr[0][:70]!r}'
            txt = sh.text_frame.text[:60].replace('\n', ' | ') if sh.has_text_frame and sh.text_frame.text.strip() else ''
            print(f'  {sh.shape_id:>3} {str(sh.shape_type):22s} {box:26s} {sh.name[:28]!r}{extra} {txt!r}' if txt
                  else f'  {sh.shape_id:>3} {str(sh.shape_type):22s} {box:26s} {sh.name[:28]!r}{extra}')


if __name__ == '__main__':
    main()
